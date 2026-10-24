"""Builds the BCSE302P mini-project report (.docx) from the official template and the REAL execution evidence
produced by `python manage.py setup` and `python manage.py uitests` (outputs/*.txt, outputs/ui_tests.json, outputs/screenshots/*.png).
usage: python build_report.py <evidence_dir> <template.docx> <out.docx>"""
import sys, os, re, json, copy
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
import report_lib as L

EV = sys.argv[1]; TPL = sys.argv[2]; OUT = sys.argv[3]
PROJ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SQL = os.path.join(PROJ, "sql"); APP = os.path.join(PROJ, "app")
WORK = os.path.join(os.path.dirname(OUT) or ".", "_img"); os.makedirs(WORK, exist_ok=True)

TITLE = "RetainAI – Customer Retention Analytics Platform"
NAME = "P R Ashwath Raaghavan"; REG = "24BCE2770"
env = json.load(open(os.path.join(EV, "environment.json")))
ui = json.load(open(os.path.join(EV, "ui_tests.json")))
sqlver = open(os.path.join(EV, "sqlplus_version.txt")).read().strip().split("\n")


def rd(path):
    return open(path, encoding="utf-8", errors="replace").read()


# ------------------------------------------------------------------ SQL tests (genuine rows of TEST_RESULTS)
def load_psv():
    rows, cur = [], None
    for ln in rd(os.path.join(EV, "test_results.psv")).split("\n"):
        if re.match(r"^TC\d\d\|", ln):
            if cur: rows.append(cur)
            cur = ln
        elif cur is not None and ln.strip():
            cur += "\n" + ln
    if cur: rows.append(cur)
    out = []
    for r in rows:
        parts = r.split("|")
        status = parts[-1].strip()
        tc, desc, inp, exp = parts[0], parts[1], parts[2], parts[3]
        act = "|".join(parts[4:-1]).split("\n")[0]
        out.append((tc, desc, inp, exp, act, status))
    return out


tests = load_psv()
n_sql = len(tests); n_sql_pass = sum(1 for t in tests if t[5] == "PASS")
n_ui = len(ui); n_ui_pass = sum(1 for t in ui if t["status"] == "PASS")
assert n_sql == 34 and n_sql_pass == 34, (n_sql, n_sql_pass)
assert n_ui == 19 and n_ui_pass == 19, (n_ui, n_ui_pass)



# ------------------------------------------------------------------ run-dependent facts, read from the evidence files
def _num(t): return float(re.sub(r"[^0-9.]", "", t))
_facts = json.load(open(os.path.join(EV, "ui_facts.json")))
_kp = {k.upper(): v for k, v in _facts["kpis"].items()}
MRR_TXT = _kp["MONTHLY RECURRING REVENUE"][0]                       # e.g. $311,917.81
RAR_TXT = _kp["REVENUE AT RISK / MONTH"][0]
HIGH_N = re.search(r"([\d,]+) active HIGH", _kp["REVENUE AT RISK / MONTH"][1]).group(1)
RAR_PCT = "%.1f" % (100 * _num(RAR_TXT) / _num(MRR_TXT))
_mlr = {}
for _r in _facts["ml_rows"]:
    _nm = "Logistic Regression" if _r.startswith("Logistic") else "Gradient Boosting"
    _mlr[_nm] = [float(x) for x in re.findall(r"\b0\.\d{3}\b", _r)][:5]   # auc, acc, prec, rec, f1
LR, GB = _mlr["Logistic Regression"], _mlr["Gradient Boosting"]
assert any(r.startswith("Logistic") and "selected" in r for r in _facts["ml_rows"]), "report text assumes Logistic Regression was selected"
_r3txt = rd(os.path.join(EV, "reports.txt")).split("=== R3.")[1].split("=== R4.")[0]
_r3 = {m.group(1): m.group(2) for m in re.finditer(r"^(HIGH|MEDIUM|LOW)\s+\d+\s+\d+\s+([\d.]+)", _r3txt, re.M)}
R3 = {k: "%.1f" % float(v) for k, v in _r3.items()}
BILL_ID = re.search(r"BILL_ID BILL_MONTH.*?\n.*?\n\s*(\d+) ", rd(os.path.join(EV, "routines.txt")).split("=== P2.")[1].split("=== P3.")[0], re.S).group(1)
_t27 = [t for t in tests if t[0] == "TC27"][0][4]
TC27_MRR = re.search(r"MRR=([\d.]+)", _t27).group(1); TC27_MRR = "{:,.2f}".format(float(TC27_MRR))
_v1 = rd(os.path.join(EV, "verify.txt")).split("=== V2.")[0]
def _vrow(t): return int(re.search(r"^%s\s+(\d+)" % t, _v1, re.M).group(1))
BILLS_TOTAL = "{:,}".format(_vrow("BILLING_RECORDS")); TICKETS_N = "{:,}".format(_vrow("SUPPORT_TICKETS"))
def _tc(n): return [t for t in tests if t[0] == n][0][4]
BILLS_N = "{:,}".format(int(re.search(r"first run=(\d+)", _tc("TC16")).group(1)))
CH_M2M = "%.1f" % float(re.search(r"package=([\d.]+)", _tc("TC25")).group(1))
CH_ALL = re.search(r"package=([\d.]+)", _tc("TC26")).group(1)
BULK_N = "{:,}".format(int(re.search(r"BULK_INSERT\s+(\d+) bills", rd(os.path.join(EV, "routines.txt"))).group(1)))
N_SHOTS = len(os.listdir(os.path.join(EV, "screenshots"))) + 7

# ------------------------------------------------------------------ spool-file helpers
def spool_block(fname, start, end=None, maxlines=None):
    lines = rd(os.path.join(EV, fname)).split("\n")
    s = next(i for i, l in enumerate(lines) if l.startswith(start) or start in l)
    e = len(lines)
    if end:
        e = next((i for i in range(s + 1, len(lines)) if end in lines[i]), len(lines))
    blk = L.squeeze(lines[s:e])
    return blk[:maxlines] if maxlines else blk


_img_no = [0]


def spool_image(fname, start, end=None, label="", maxlines=None, drop_prefix=None):
    blk = spool_block(fname, start, end, maxlines)
    if drop_prefix:
        blk = [l for l in blk if not l.startswith(drop_prefix)]
    _img_no[0] += 1
    out = os.path.join(WORK, f"spool_{_img_no[0]:02d}.png")
    L.render_text(blk, out, f"Rendered from the real SQL*Plus spool file outputs\\{fname}  |  user RETAINAI @ XEPDB1", size=15)
    return out


def obj_source(text, kind, name):
    m = re.search(r"CREATE OR REPLACE %s\s+%s\b.*?\n/\n" % (kind, name), text, re.S | re.I)
    return m.group(0).rstrip("\n") if m else ""


# ------------------------------------------------------------------ document
doc = Document(TPL)
body = doc.element.body
kids = list(body.iterchildren())

# styles: outline levels so that the TOC field can collect the headings
for sid, lvl in (("Heading 1", 0), ("Heading 2", 1)):
    st = next(x for x in doc.styles if x.name == sid)
    ppr = st.element.get_or_add_pPr()
    for old in ppr.findall(qn("w:outlineLvl")): ppr.remove(old)
    o = OxmlElement("w:outlineLvl"); o.set(qn("w:val"), str(lvl)); ppr.append(o)


def set_para_text(el, text, bold=False, italic=False, size=12, align=None):
    p = Paragraph(el, doc)
    for r in list(p.runs): r._element.getparent().remove(r._element)
    L.add_runs(p, text, size, bold, italic)
    return p


# cover page ------------------------------------------------------------------
set_para_text(kids[6], TITLE, bold=True, size=18)
for r in Paragraph(kids[6], doc).runs: r.font.color.rgb = __import__("docx").shared.RGBColor.from_string("1F3864")
ct = Table(kids[13], doc)
for row in list(ct.rows)[2:]:
    ct._tbl.remove(row._tr)
for i, v in enumerate(["1", NAME, REG]):
    c = ct.rows[1].cells[i]; p = c.paragraphs[0]
    for r in list(p.runs): r._element.getparent().remove(r._element)
    run = p.add_run(v); L._fonts(run, L.TNR, 11)
import copy as _cp
for n in ("2", "3"):                       # two empty rows for additional team members
    tr = _cp.deepcopy(ct.rows[1]._tr); ct._tbl.append(tr)
    row = ct.rows[-1]
    for i, v in enumerate([n, "", ""]):
        p = row.cells[i].paragraphs[0]
        for r in list(p.runs): r._element.getparent().remove(r._element)
        if v:
            run = p.add_run(v); L._fonts(run, L.TNR, 11)
body.remove(kids[14])
set_para_text(kids[21], "Date of Submission: 21 / 10 / 2026", bold=True)
Paragraph(kids[21], doc).alignment = L.AL["center"]

# certificate -------------------------------------------------------------------
cert = set_para_text(kids[24],
    f"This is to certify that the Mini Project Report titled “**{TITLE}**” submitted by **{NAME} (Reg. No. {REG})** "
    "in partial fulfilment of the requirements of the course **BCSE302P – Database Systems Laboratory** at the School of Computer "
    "Science and Engineering (SCOPE), Vellore Institute of Technology, Vellore, is a bona fide record of the work carried out "
    "under my guidance during the Fall Semester 2026–27.")
cert.alignment = L.AL["both"]

# declaration block, moved behind the faculty block -------------------------------------
decl_start = len(list(body.iterchildren())) - 1
L.heading(doc, "DECLARATION", 1, page_break=True)
L.para(doc, f"I, **{NAME}** (Reg. No. **{REG}**), declare that this mini project, “{TITLE}”, is my own work, carried out "
            "individually. The database design, SQL and PL/SQL code, web application, test cases and this report were prepared for the "
            "course BCSE302P and have not been submitted elsewhere.")
L.para(doc, "Every execution result, test outcome and screenshot in this report comes from running the project's own scripts on my laptop (Oracle Database 21c Express Edition). "
            "Nothing has been invented. Long SQL outputs are shown as images rendered from the real spool files, the terminal screenshots were saved from the Visual Studio Code window while the commands ran, "
            "and a few very long browser screenshots are cropped to their upper part. The limits of the data and of the test setup are stated in Sections 3.5, 6.3 and 7.2. "
            "Public sources are listed in the References.")
L.para(doc, "", after=24)
L.para(doc, f"Place: Vellore", align="left", after=0)
L.para(doc, "Date: 21 / 10 / 2026", align="left", after=18)
L.para(doc, f"Signature: ______________________      **{NAME}**", align="left")
decl_els = list(body.iterchildren())[decl_start:-1]
set_para_text(kids[35], "Date: ____ / ____ / 2026")
anchor = kids[35]
for el in decl_els:
    anchor.addnext(el); anchor = el

# abstract ---------------------------------------------------------------------------
abstract = (
    "Retention teams at telecom operators tend to work from spreadsheets, and a spreadsheet cannot say which customers are about to leave or whether last month's offer did any good. "
    "RetainAI tackles this on Oracle Database 21c Express Edition. The public IBM/Kaggle Telco Customer Churn data (7,043 customers) goes into a staging table, and a PL/SQL procedure turns it "
    "into a third-normal-form schema of 17 tables. Most of the logic sits in the database itself: seven functions, among them a rule-based churn-risk score and a customer segmentation, "
    "twenty procedures, a package that hands result sets to the front end as REF cursors, ten triggers, and all five kinds of cursor. A responsive Flask web application, connected through "
    "python-oracledb, covers customers, billing, support tickets, retention campaigns, analytics and administration. It also offers an optional scikit-learn churn model whose probabilities are "
    "written back to Oracle. In the final run on the author's laptop all 34 SQL/PL-SQL tests and all 19 browser tests passed. Names, billing history, tickets and campaigns are synthetic and are marked as such.")
wc = len(abstract.split()); assert 150 <= wc <= 200, wc
set_para_text(kids[37], abstract).alignment = L.AL["both"]
kp = Paragraph(kids[38], doc)
for r in list(kp.runs): r._element.getparent().remove(r._element)
L.add_runs(kp, "**Keywords:** DBMS, Oracle PL/SQL, Customer Churn, Normalization, Triggers, Flask")
body.remove(kids[41])

# remove template body from chapter 1 onwards ----------------------------------------------
cur = list(body.iterchildren())
start_idx = next(i for i, e in enumerate(cur) if e is kids[42])
for e in cur[start_idx:-1]:
    body.remove(e)

# =====================================================================================
# 1. INTRODUCTION
# =====================================================================================
L.heading(doc, "1. INTRODUCTION", 1, page_break=True)
L.heading(doc, "1.1 Problem Statement", 2)
L.para(doc, "Telecom operators lose customers all the time (the term is “churn”), and signing up a new customer costs far more than keeping an old one. "
            "In the dataset used here 1,869 of 7,043 customers (26.5%) have left. The split by contract is stark: 42.7% of month-to-month customers left, against 2.8% of those on two-year contracts. "
            "A retention team needs four answers: **who** is likely to leave, **why**, **how much revenue** that puts at risk, and **whether the offers it sends make any difference**.")
L.para(doc, "The people who would use such a system are retention managers, support agents, billing staff and an administrator. In practice they often work from exported spreadsheets. "
            "The Telco source file is one flat table of 21 columns in which contract names, payment methods and service flags repeat 7,043 times. Nothing stops invalid or duplicate rows. "
            "Billing, tickets and retention actions sit in separate files that cannot be joined reliably, and nothing records who changed what. "
            "RetainAI moves all of this into a normalised Oracle database, so the integrity rules, the risk scoring and the audit trail are enforced by the database rather than by habit.")
L.heading(doc, "1.2 Objectives", 2)
L.numbered(doc, [
    "Design a normalised (3NF) schema for customers, subscriptions, billing, support tickets, retention campaigns and churn scores, and load the real Telco data into it through an ETL procedure.",
    "Put the business rules inside the database as PL/SQL: functions for risk score, segment, tenure and lifetime value; procedures for registration, cancellation, billing, tickets and campaigns; a package for analytics; triggers for validation and auditing; and cursors of each kind.",
    "Produce churn analytics, customer segments, a ranked retention list and campaign-effectiveness reports with advanced SQL (window functions, ROLLUP, anti-joins).",
    "Build a responsive web application that does everything through the stored procedures and draws its charts from live Oracle queries.",
    "Add an optional machine-learning churn model that reads its training data from Oracle and writes its probabilities back.",
    "Test everything with repeatable SQL and browser tests, and record the results instead of asserting them.",
])
L.heading(doc, "1.3 Importance of DBMS", 2)
L.para(doc, "A DBMS is the right foundation here because a retention system is only as good as the data it shares.")
L.bullets(doc, [
    "**Integrity.** Primary, foreign, unique and check constraints turn bad data away at the door. A function-based unique index makes sure a customer has at most one ACTIVE subscription even when two users act at the same moment, and triggers stop a paid bill from being reopened.",
    "**Less redundancy.** Contract types, payment methods, plans and services are stored once and referenced by key (Section 3.4), so a change is made in a single row.",
    "**Concurrency.** Oracle’s read consistency lets the dashboard keep reading while agents register customers or record payments.",
    "**Security.** The web application never builds business state on its own. It calls stored procedures that validate every input, and the privileges belong to a dedicated schema owner.",
    "**Atomicity.** A procedure either commits all of its work or rolls it back. Cancelling a subscription, for example, also marks the customer CHURNED and queues a win-back action, and these happen together or not at all.",
    "**Speed.** Indexes and set-based SQL work out churn rates, rankings and revenue at risk for 7,043 customers and "+BILLS_TOTAL+" bills inside the database, with no loops in the application.",
])
L.heading(doc, "1.4 Scope of the Project", 2)
L.para(doc, "**Included:** the Oracle schema, ETL and seed data; seven functions, twenty procedures, one package and ten triggers; every cursor type; "
            "and a Flask web application with login and password recovery, a dashboard, customer management, billing, tickets, retention campaigns, analytics, reports with CSV export, "
            "an ML module and an administration page. The SQL and browser tests are automated.")
L.para(doc, "**Not included:** a real payment gateway, or sending e-mail and SMS (a retention “action” is a database record, not a message). The web layer has a single administrator login "
            "and no per-user roles. There is no multi-tenant operation, no load or penetration testing and no deployment to a server. The data are a historical snapshot, not a live feed.")

# =====================================================================================
# 2. SYSTEM REQUIREMENTS
# =====================================================================================
L.heading(doc, "2. SYSTEM REQUIREMENTS", 1, page_break=True)
L.para(doc, "Everything below was read from the laptop that produced the results in this report; `scripts/collect_env.py` collects it automatically into `outputs/environment.json`.")
L.heading(doc, "2.1 Hardware Requirements", 2)
L.table_caption(doc, "Table 2.1:", "Hardware used for development, execution and testing")
L.table(doc, ["Component", "Specification"], [
    ["Processor", env["cpu"]], ["Memory (RAM)", f"{env['ram_gb']} GB"], ["Storage", env["disk"] + " (solid-state)"],
    ["Display / input", "Laptop screen, keyboard and mouse (browser tests also ran headless at 1366×900 and 390×844)"],
], [1.8, 4.4], size=11)
L.heading(doc, "2.2 Software Requirements", 2)
L.table_caption(doc, "Table 2.2:", "Software used")
L.table(doc, ["Category", "Tool / Version"], [
    ["Operating system", env["os"]],
    ["DBMS", "Oracle Database 21c Express Edition, Release 21.3.0.0.0 (pluggable database XEPDB1)"],
    ["SQL client", f"{sqlver[0]} {sqlver[1]}"],
    ["Language / runtime", f"Python {env['python']}"],
    ["Web framework", f"Flask {env['flask']} (Jinja2 templates), HTML5, CSS3, hand-written SVG charts (no external JavaScript library)"],
    ["Database driver", f"python-oracledb {env['oracledb']} (thin mode, no Oracle client required)"],
    ["Machine learning", f"scikit-learn {env['sklearn']}, pandas {env['pandas']}, NumPy {env['numpy']}"],
    ["Browser and UI testing", f"Microsoft Edge {env['edge']} driven by Playwright (Python)"],
    ["Editor", "Visual Studio Code; the project is run from its integrated terminal with `python manage.py`"],
    ["Diagrams", "Graphviz (ER diagram)"],
], [1.8, 4.4], size=11)
L.para(doc, "**Dataset:** “Telco Customer Churn” (IBM sample data published on Kaggle), 7,043 rows × 21 columns, stored in `data/telco.csv`.")

# =====================================================================================
# 3. DATABASE DESIGN
# =====================================================================================
L.heading(doc, "3. DATABASE DESIGN", 1, page_break=True)
L.heading(doc, "3.1 Entity–Relationship (ER) Diagram", 2)
L.para(doc, "The model has twelve core entities. A **customer** can have several **subscriptions** over time, but only one of them ACTIVE at once. Each subscription points to a **plan**, a **contract type** "
            "and a **payment method**, owns many **billing records** and can include many add-on **services**; that many-to-many link is resolved by the associative table `subscription_services`. "
            "Customers also raise **support tickets**, receive **retention actions** (which may belong to a **retention campaign**) and have exactly one current **churn score**. "
            "Five supporting tables (`app_config`, `stg_telco`, `ml_model_runs`, `audit_log`, `test_results`) hold the configuration, the raw CSV rows, the ML run history, "
            "the trigger-written audit trail and the test results.")
L.table_caption(doc, "Table 3.0:", "Relationships and cardinalities")
L.table(doc, ["Parent (1)", "Child (N)", "Cardinality", "Participation / rule"], [
    ["customers", "subscriptions", "1 : N", "A customer may have many subscriptions over time; a unique function-based index allows only one ACTIVE"],
    ["plans / contract_types / payment_methods", "subscriptions", "1 : N", "Every subscription has exactly one of each (mandatory FK)"],
    ["subscriptions", "billing_records", "1 : N", "One bill per subscription per month (UNIQUE subscription_id, bill_month)"],
    ["subscriptions ↔ services", "subscription_services", "M : N", "Associative table with composite primary key"],
    ["customers", "support_tickets", "1 : N", "A customer may raise any number of tickets"],
    ["customers", "retention_actions", "1 : N", "Calls, e-mails, discounts, upgrades and win-back actions"],
    ["retention_campaigns", "retention_actions", "1 : N (optional)", "An action may be ad hoc (campaign_id NULL) – for example one created by a trigger"],
    ["customers", "churn_scores", "1 : 0..1", "Current rule-based score, risk band, segment and optional ML probability"],
], [1.7, 1.4, 0.9, 2.2], size=10)
L.picture(doc, os.path.join(PROJ, "docs", "img", "er_diagram.png"), 6.3)
L.caption(doc, f"Figure 3.1: ER Diagram of {TITLE} (crow’s-foot notation; generated from docs/er.dot with Graphviz)")

# ---- 3.2 relational schema (generated from the DDL itself, so it cannot drift from the real database)
L.heading(doc, "3.2 Relational Schema", 2)
L.para(doc, "Every relation below is generated from `sql/01_schema.sql`, the script that created the tables in Oracle, so the listing cannot drift from the real database. "
            "PK is primary key, FK foreign key, UQ unique and NN not null.")
DESC = {
 "app_config": ("Business configuration; REFERENCE_DATE is the snapshot date used as “today”.", {}),
 "stg_telco": ("Staging copy of the raw CSV (all text), loaded by 02_load_staging.sql and consumed by the ETL procedure.", {}),
 "contract_types": ("Lookup of contract terms.", {"term_months": "Commitment length (1, 12, 24)"}),
 "payment_methods": ("Lookup of payment methods.", {"is_automatic": "Y for automatic payment"}),
 "plans": ("Service plans derived from the internet/phone combination.", {"base_price": "Lowest monthly price seen for the plan"}),
 "services": ("Add-on services (Online Security, Tech Support, …).", {"service_category": "Security, Support, Streaming, …"}),
 "customers": ("One row per customer; deleting is blocked by a trigger.", {"status": "ACTIVE or CHURNED", "email": "Normalised to lower case by trigger"}),
 "subscriptions": ("Contract of a customer with a plan, contract type and payment method.", {"monthly_charge": "Charge billed every month", "cancel_reason": "Filled when cancelled"}),
 "subscription_services": ("Associative table: add-ons chosen for a subscription.", {}),
 "billing_records": ("Monthly bills (synthetic history built from tenure and charge).", {"bill_month": "First day of the billed month", "due_date": "Last day of the month by default"}),
 "support_tickets": ("Support requests (synthetic).", {"category": "BILLING, TECHNICAL, SERVICE, CANCELLATION, OTHER"}),
 "retention_campaigns": ("Retention offers targeted at a segment.", {"discount_pct": "Discount applied when accepted"}),
 "retention_actions": ("Single contact/offer made to a customer.", {"outcome": "PENDING, ACCEPTED, DECLINED, NO_RESPONSE"}),
 "churn_scores": ("Current score of every customer.", {"rule_score": "0–100 from fn_churn_risk_score", "ml_probability": "Churn probability from the ML model"}),
 "ml_model_runs": ("History of ML training runs and metrics.", {}),
 "audit_log": ("Audit trail written only by triggers.", {"changed_by": "Session user"}),
 "test_results": ("Results of the automated SQL tests.", {}),
}


def parse_schema(sql):
    tables = []
    for m in re.finditer(r"CREATE TABLE (\w+) \((.*?)\n\);", sql, re.S):
        name, body_ = m.group(1), m.group(2)
        parts, depth, cur_ = [], 0, ""
        for ch in body_:
            if ch == "(": depth += 1
            if ch == ")": depth -= 1
            if ch == "," and depth == 0:
                parts.append(cur_); cur_ = ""
            else:
                cur_ += ch
        parts.append(cur_)
        cols, tcons = [], []
        for p in parts:
            p = " ".join(p.split())
            if not p: continue
            if p.upper().startswith(("CONSTRAINT ", "PRIMARY KEY", "UNIQUE")) and not re.match(r"^\w+ (VARCHAR2|NUMBER|CHAR|DATE|TIMESTAMP)", p):
                tcons.append(re.sub(r"^CONSTRAINT \w+ ", "", p)); continue
            m2 = re.match(r"^(\w+) (VARCHAR2\(\d+\)|NUMBER(?:\(\d+(?:,\d+)?\))?|CHAR\(\d+\)|DATE|TIMESTAMP)(.*)$", p)
            if not m2: continue
            col, typ, rest = m2.groups()
            cons = []
            if "PRIMARY KEY" in rest: cons.append("PK")
            if "GENERATED" in rest: cons.append("identity")
            fk = re.search(r"REFERENCES (\w+) \((\w+)\)", rest)
            if fk: cons.append(f"FK → {fk.group(1)}({fk.group(2)})")
            if re.search(r"\bUNIQUE\b", rest): cons.append("UQ")
            ck = re.search(r"CHECK \((.*)\)", rest)
            if ck: cons.append("CHECK " + ck.group(1)[:70])
            if "NOT NULL" in rest: cons.append("NN")
            d = re.search(r"DEFAULT (.+?)(?: NOT NULL| CONSTRAINT|$)", rest)
            if d: cons.append("DEFAULT " + d.group(1).strip())
            cols.append((col, typ, ", ".join(cons) or "–"))
        tables.append((name, cols, tcons))
    return tables


ddl_text = rd(os.path.join(SQL, "01_schema.sql"))
schema = parse_schema(ddl_text)
assert len(schema) == 17, len(schema)
for k, (name, cols, tcons) in enumerate(schema, 1):
    desc, colnotes = DESC[name]
    L.table_caption(doc, f"Table 3.{k}:", f"{name.upper()} – {desc}")
    L.table(doc, ["Attribute", "Data type", "Constraints", "Note"],
            [[c, t, k_, colnotes.get(c, "")] for c, t, k_ in cols], [1.35, 1.1, 2.55, 1.2], size=8.5, code_cols=(0,))
    if tcons:
        L.para(doc, "Table-level constraints: " + "; ".join("`%s`" % t[:110] for t in tcons) + ".", size=10, after=8)

# ---- 3.3 DDL
L.heading(doc, "3.3 Schema Definition (DDL)", 2)
L.para(doc, "The full script `sql/01_schema.sql` follows (17 tables, 1 sequence and 9 indexes). It ran unchanged on Oracle; Figures 3.2 and 3.3 show the outcome.")
L.code_block(doc, "\n".join(l for l in ddl_text.split("\n") if not l.startswith("SET ")), size=8)
L.picture(doc, spool_image("verify.txt", "=== V1.", "=== V3.", maxlines=40), 5.6)
L.caption(doc, "Figure 3.2: Row counts of all tables and status of the schema objects after the build (real output, rendered from the spool file)")
L.picture(doc, spool_image("verify.txt", "=== V5.", "=== V6."), 3.6)
L.caption(doc, "Figure 3.3: Integrity constraints in the schema (real output: 30 CHECK, 16 PRIMARY KEY, 11 FOREIGN KEY, 7 UNIQUE)")

# ---- 3.4 normalization
L.heading(doc, "3.4 Normalization", 2)
L.para(doc, "**Unnormalised form (UNF).** The source is a single table, `TELCO(customerID, gender, SeniorCitizen, Partner, Dependents, tenure, "
            "PhoneService, MultipleLines, InternetService, OnlineSecurity, OnlineBackup, DeviceProtection, TechSupport, StreamingTV, StreamingMovies, "
            "Contract, PaperlessBilling, PaymentMethod, MonthlyCharges, TotalCharges, Churn)`. It has a repeating group (six add-on columns), "
            "text flags (Yes, No, “No internet service”), `TotalCharges` stored as text with some blanks, and facts about contracts and payment "
            "methods repeated thousands of times.")
L.para(doc, "**Functional dependencies found in the data:**", keep_next=True)
L.bullets(doc, [
    "customerID → gender, SeniorCitizen, Partner, Dependents, tenure, PaperlessBilling, MonthlyCharges, TotalCharges, Churn, Contract, PaymentMethod, InternetService, PhoneService",
    "Contract → term_months (Month-to-month = 1, One year = 12, Two year = 24)",
    "PaymentMethod → is_automatic (the method name ends with “(automatic)”)",
    "(InternetService, PhoneService) → plan, base price",
    "(customerID, service) → chosen / not chosen (a multi-valued fact, one per add-on)",
    "customerID → email (candidate key, unique)",
])
L.table_caption(doc, "Table 3.18:", "Step-by-step decomposition")
L.table(doc, ["Normal form", "Violation found in the Telco table", "Decomposition applied"], [
    ["1NF", "Six add-on columns form a repeating group; values such as “No internet service” mix a flag and a fact; TotalCharges is text.",
     "Add-ons become rows of `subscription_services(subscription_id, service_id)`; flags become CHAR(1) Y/N; numbers are converted by `fn_to_num` in the ETL."],
    ["2NF", "With the composite key (customerID, service) the attribute `service_category` depends on service only (partial dependency).",
     "`services(service_id, service_name, service_category)` is a separate relation referenced by the associative table."],
    ["3NF", "Transitive dependencies: customerID → Contract → term_months; customerID → PaymentMethod → is_automatic; customerID → (Internet, Phone) → plan, price.",
     "New relations `contract_types`, `payment_methods`, `plans`; `subscriptions` holds only foreign keys to them."],
    ["BCNF", "Every determinant of every relation is now a candidate key (customer_id, subscription_id, bill_id, ticket_id, …; the unique e-mail and (subscription_id, bill_month) are candidate keys too).",
     "No further decomposition is needed. Customer attributes and subscription attributes are also separated so that a customer can have several subscriptions over time."],
], [0.7, 3.0, 2.5], size=9.5)
L.para(doc, "The decomposition is lossless. `sp_load_from_staging` rebuilds every original row from the new tables, and Figure 3.4 shows that the 7,043 staged rows "
            "produced 7,043 customers with the same 1,869 churned.")
L.picture(doc, spool_image("verify.txt", "=== V7."), 4.4)
L.caption(doc, "Figure 3.4: Source check after the ETL – staged rows equal loaded customers and churned customers (real output)")

# ---- 3.5 data honesty
L.heading(doc, "3.5 Data Sources: What Is Real and What Is Synthetic", 2)
L.para(doc, "The table separates what comes from the public dataset from what the project generated, so that nothing synthetic is mistaken for real data.")
L.table_caption(doc, "Table 3.19:", "Provenance of the data")
L.table(doc, ["Data", "Origin", "How it was produced"], [
    ["Demographics, contract, internet/phone service, add-ons, payment method, monthly charge, tenure, churn label", "Real (public Telco dataset, 7,043 rows)", "Loaded unchanged from `data/telco.csv`"],
    ["Customer full name and e-mail", "Synthetic", "`scripts/gen_staging.py`, fixed random seed 2770; the dataset has only anonymised IDs"],
    ["Subscription start/end dates", "Derived", "From tenure and the business reference date 30-Sep-2026 (the dataset has no calendar dates)"],
    ["Billing history ("+BILLS_TOTAL+" bills)", "Synthetic", "One bill per month of tenure at the real monthly charge; paid, pending or overdue status assigned by rules in the ETL"],
    ["Support tickets ("+TICKETS_N+")", "Synthetic", "`sp_seed_support_tickets`; churned customers are given a higher chance of having tickets (40% against 18%), so any ticket–churn relation is built in"],
    ["Retention campaigns, actions and responses", "Synthetic / simulated", "`sp_seed_campaigns`, `sp_plan_retention`, and simulated responses marked “Simulated response (demo data)”"],
    ["Risk score, segment, ML probability", "Computed", "By `fn_churn_risk_score`, `fn_segment` and the scikit-learn model"],
], [2.3, 1.5, 2.4], size=9.5)

# =====================================================================================
# 4. IMPLEMENTATION
# =====================================================================================
L.heading(doc, "4. IMPLEMENTATION", 1, page_break=True)
L.para(doc, "Every object in this chapter was created by the scripts in `sql/` and run on Oracle. The figures are real SQL*Plus output captured with `SPOOL` and rendered as images, "
            "and each caption names its spool file. Screenshots of the same commands running in the Visual Studio Code terminal are in Section 6.2. "
            "The demonstrations come from `sql/14_demo_routines.sql` (functions, procedures, triggers) and `sql/10_cursors_demo.sql` (cursors).")

fig4 = [0]


def fig(img, text, width=6.0):
    fig4[0] += 1
    L.picture(doc, img, width)
    L.caption(doc, f"Figure 4.{fig4[0]}: {text}")


# ---- 4.1 functions
fsrc = rd(os.path.join(SQL, "03_functions.sql"))
fsrc_clean = "\n".join(l for l in fsrc.split("\n") if not l.startswith(("SET ", "SELECT object_name")))
L.heading(doc, "4.1 Functions", 2)
L.para(doc, "Seven stored functions hold the reusable logic. `fn_to_num` converts text to numbers regardless of the session locale and is used by the ETL, and `fn_ref_date` returns the business date "
            "from `app_config`. `fn_tenure_months` and `fn_lifetime_value` measure how long a customer has stayed and what they have paid. `fn_risk_band` maps a score to LOW, MEDIUM or HIGH. "
            "`fn_churn_risk_score` is the rule-based model: it adds points for the churn drivers seen in the data (month-to-month contract +30, a new customer up to +20, fibre optic +10, electronic check +10, "
            "no Tech Support or Online Security, unresolved urgent tickets, overdue bills) and caps the total at 100. `fn_segment` returns Lost, New, At Risk, Champion, Loyal or Standard. "
            "Procedures, views, triggers and the web application all call them.")
L.code_block(doc, fsrc_clean, size=8)
fig(spool_image("routines.txt", "=== F1.", "=== P1."), "The functions applied to six sample customers (real output, spool file routines.txt)", 5.8)
L.para(doc, "Test cases TC09–TC12 (Table 6.1) confirm that a risky profile scores 88 and a loyal one 17, that the band boundaries 34, 35, 59 and 60 map to LOW, MEDIUM, MEDIUM and HIGH, "
            "that customer 5575-GNVDE has a tenure of 34 months and a lifetime value of 1,936.30, and that a churned customer is segmented as Lost.")

# ---- 4.2 procedures
psrc = rd(os.path.join(SQL, "06_procedures.sql"))
L.heading(doc, "4.2 Procedures", 2)
L.para(doc, "Twenty procedures carry out the operations of the application, and the web screens call nothing else to change data. Each one validates its input, raises a numbered business error "
            "(ORA-20xxx) when a rule is broken, and commits only when all of its work succeeded. The seven below matter most; the complete source is in Appendix A.")
for nm, why in [
    ("sp_add_customer", "Registers a customer, creates the first subscription and its add-ons and scores the customer (IN: personal data, plan, contract, payment, add-on list; OUT: new customer id and subscription id)."),
    ("sp_cancel_subscription", "Cancels the subscription with a reason; a trigger then marks the customer CHURNED and queues a win-back action."),
    ("sp_generate_bills", "Month-end billing run: inserts one bill per active subscription (OUT: number created); running it twice creates nothing the second time."),
    ("sp_record_payment", "Marks a bill PAID; raises ORA-20004 if the bill is missing or already paid."),
    ("sp_open_ticket", "Opens a support ticket (OUT: ticket id); an urgent cancellation ticket makes a trigger queue a retention call."),
    ("sp_plan_retention", "Uses a parameterised cursor to create one retention action for every ACTIVE HIGH-risk customer of a campaign’s target segment."),
    ("sp_record_outcome", "Records the customer’s answer; an accepted discount offer reduces the monthly charge exactly once."),
]:
    L.subheading(doc, nm)
    L.para(doc, why, after=3)
    L.code_block(doc, obj_source(psrc, "PROCEDURE", nm), size=8)
fig(spool_image("routines.txt", "=== P1.", "=== P2."), "sp_add_customer: a customer is registered; the trigger normalised the name and e-mail; the score row was created (real output)", 6.0)
fig(spool_image("routines.txt", "=== P2.", "=== P3."), "sp_generate_bills and sp_record_payment (real output). Bill "+BILL_ID+" was created by the billing run and then paid", 5.4)
fig(spool_image("routines.txt", "=== P3.", "=== P4."), "sp_open_ticket and the retention call queued automatically by trigger trg_tickets_ai (real output)", 6.0)
fig(spool_image("routines.txt", "=== P4.", "=== T2."), "sp_cancel_subscription: the customer becomes CHURNED, the segment becomes Lost and a win-back action appears (real output)", 6.0)

# ---- 4.3 package
pk = rd(os.path.join(SQL, "07_package.sql"))
pk_clean = "\n".join(l for l in pk.split("\n") if not l.startswith(("SET ", "SHOW ERRORS")))
L.heading(doc, "4.3 Packages", 2)
L.para(doc, "`pkg_retainai` is the analytics interface of the application. Its specification exposes scalar KPIs (`churn_rate`, `monthly_recurring_revenue`, `revenue_at_risk`, "
            "`customers_in_band`), result sets returned as `SYS_REFCURSOR` (`dashboard_kpis`, `segment_report`, `churn_breakdown`, `high_risk_report`) and the batch procedure `refresh_analytics`. "
            "The dashboard, analytics and reports pages all call it. Keeping the logic in the package means SQL*Plus, the web application or any other client gets the same figures.")
L.code_block(doc, pk_clean, size=8)
cur_blk = spool_block("cursors.txt", "Demo 6")
fig(spool_image("cursors.txt", "Demo 5", "Demo 6", maxlines=14), "Package procedure pkg_retainai.segment_report returning a REF CURSOR (real output, spool file cursors.txt)", 5.8)
fig(spool_image("cursors.txt", "Demo 6"), "Package procedure pkg_retainai.dashboard_kpis printed through a REF CURSOR bind variable (real output)", 5.8)
L.para(doc, "Test cases TC25–TC29 compare the package with independent SQL calculations ("+CH_M2M+"% month-to-month churn, "+CH_ALL+"% overall, MRR "+TC27_MRR+") and check "
            "that an unknown dimension raises ORA-20010.")

# ---- 4.4 triggers
tg = rd(os.path.join(SQL, "08_triggers.sql"))
tg_clean = "\n".join(l for l in tg.split("\n") if not l.startswith(("SET ", "SELECT trigger_name")))
L.heading(doc, "4.4 Triggers", 2)
L.para(doc, "Ten triggers enforce rules that must hold whichever program touches the data. Four fire BEFORE a row is written: `trg_customers_bi` generates the id, lower-cases the e-mail, "
            "applies INITCAP to the name and validates the address, while `trg_billing_biu`, `trg_tickets_biu` and `trg_actions_bi` guard their own tables. "
            "`trg_customers_bd` fires BEFORE DELETE and forbids deleting customers. Four AFTER row-level triggers (`trg_subscriptions_audit`, `trg_subscriptions_status`, `trg_tickets_ai`, "
            "`trg_scores_band_change`) handle auditing and cascades. The last, `trg_billing_compound`, is a COMPOUND trigger that writes one audit row for a whole bulk insert instead of thousands.")
L.code_block(doc, tg_clean, size=8)
fig(spool_image("routines.txt", "=== T1.", "=== P4.", maxlines=12), "trg_tickets_ai: an urgent cancellation ticket created a retention CALL automatically (real output)", 6.0)
fig(spool_image("routines.txt", "=== T2.", "=== T3."), "Audit trail written by triggers: subscription INSERT/UPDATE rows and ONE BULK_INSERT row for "+BULK_N+" bills (real output)", 6.3)
fig(spool_image("routines.txt", "=== T3.", "=== Clean-up"), "Business rules enforced by triggers and constraints; every error shown is the expected rejection (real output)", 6.3)

# ---- 4.5 cursors
cs = rd(os.path.join(SQL, "10_cursors_demo.sql"))
L.heading(doc, "4.5 Cursors", 2)
L.para(doc, "Some work has to happen row by row, because each row triggers logic that one SQL statement cannot express. Six demonstrations in `sql/10_cursors_demo.sql` cover every kind of cursor, "
            "and the production code uses the same techniques: `sp_score_all_customers` an explicit cursor, `sp_plan_retention` a parameterised one, `pkg_retainai.high_risk_report` a REF cursor.")
L.code_block(doc, "\n".join(l for l in cs.split("\n") if not l.startswith(("SET ", "SPOOL", "COLUMN"))), size=8)
fig(spool_image("cursors.txt", "=====", "Demo 3", maxlines=22), "Implicit cursor (SQL%ROWCOUNT) and explicit cursor (OPEN/FETCH/%NOTFOUND/CLOSE) – real output", 6.0)
fig(spool_image("cursors.txt", "Demo 3", "Demo 5"), "Parameterised cursors in nested loops and BULK COLLECT ... LIMIT (revenue at risk cross-checked with the package) – real output", 6.0)

# ---- 4.6 ETL / views / reports
L.heading(doc, "4.6 ETL, Views and Analytical Queries", 2)
L.para(doc, "`sp_load_from_staging` (Appendix A) fills the lookup tables with `INSERT ... SELECT` and `ROW_NUMBER()`, then loads customers, subscriptions, add-ons and bills. "
            "Eight views (`v_customer_current`, `v_customer_360`, `v_churn_by_dimension`, `v_segment_summary`, `v_risk_distribution`, `v_monthly_revenue`, `v_ticket_summary`, "
            "`v_campaign_effectiveness`) feed the screens. `sql/11_reports.sql` holds eleven reports built from aggregates, `RANK()`, running totals with `SUM() OVER`, `ROLLUP`, "
            "outer joins and a `NOT EXISTS` anti-join. A few outputs follow.")
fig(spool_image("reports.txt", "=== R1.", "=== R3."), "Reports R1–R2: churn rate by contract (aggregate + RANK) and by tenure band, payment method and internet service (real output)", 6.3)
fig(spool_image("reports.txt", "=== R3.", "=== R5."), "Reports R3–R4: validation of the risk score (observed churn per band) and customer segmentation (real output)", 6.3)
fig(spool_image("reports.txt", "=== R9.", "=== R10."), "Report R9: active revenue by contract and payment method with ROLLUP sub-totals (real output)", 5.4)
fig(spool_image("reports.txt", "=== R10.", "=== R11."), "Report R10: fibre-optic customers without Tech Support – 49.4% churned (NOT EXISTS anti-join, real output)", 5.0)
L.para(doc, "Report R3 is the honest check on the rule-based score: of the customers scored HIGH, "+R3['HIGH']+"% did churn, against "+R3['MEDIUM']+"% of MEDIUM and "+R3['LOW']+"% of LOW, so the score does separate risk. "
            "R7 (tickets against churn) is **not** evidence of cause, because the tickets are simulated with churn built into their probability (Section 3.5).")

# =====================================================================================
# 5. FRONT-END
# =====================================================================================
L.heading(doc, "5. FRONT-END", 1, page_break=True)
L.para(doc, "The responsive web application is built with Flask, Jinja2, CSS and SVG. Every screenshot below is genuine: `scripts/capture.py` starts the application, signs in as the administrator "
            "and drives Microsoft Edge through Playwright against the live Oracle database. Long pages are cropped to their upper part.")
L.heading(doc, "5.1 User Interface", 2)
SH = os.path.join(EV, "screenshots")
f5 = [0]


def shot(fn, text, width=6.2, crop_h=None):
    f5[0] += 1
    im = Image.open(os.path.join(SH, fn))
    path = os.path.join(WORK, "s_" + fn)
    if crop_h and im.height > crop_h:
        im = im.crop((0, 0, im.width, crop_h))
    im.save(path)
    L.picture(doc, path, width)
    L.caption(doc, f"Figure 5.{f5[0]}: {text}")


shot("01_login.png", "Login page")
shot("24_forgot_password.png", "Forgot-password page: a one-time code is printed in the terminal that runs the application")
shot("02_dashboard.png", "Dashboard: KPIs and charts computed by pkg_retainai and views, retention priority list", 6.0, 1500)
L.para(doc, "The dashboard shows "+_kp["CUSTOMERS"][0]+" customers ("+_kp["CUSTOMERS"][1].replace(" · ", ", ")+"), a churn rate of "+_kp["CHURN RATE"][0]+", monthly recurring revenue of "+MRR_TXT+" and revenue at risk of "+RAR_TXT+" "
            "("+HIGH_N+" active HIGH-risk customers), with open tickets, churn by contract and by tenure band, the segment donut, the billing trend and the ranked retention list.")
shot("03_customers.png", "Customer list with search, filters and risk score", 6.0, 1300)
shot("05_customer_detail.png", "Customer detail: risk, subscription, actions (each button calls a stored procedure), billing, tickets and retention history", 6.0)
shot("06_customer_new.png", "New-customer form (calls sp_add_customer)", 6.0)
shot("07_billing.png", "Billing: month-end run, overdue marking and payments", 6.0, 1300)
shot("08_tickets.png", "Support tickets", 6.0, 1300)
shot("09_analytics.png", "Churn analytics and segmentation", 6.0)
shot("10_retention.png", "Retention campaigns, effectiveness and actions", 6.0, 1500)
shot("11_reports.png", "Reports with CSV export", 6.0, 1400)
shot("13_ml_results.png", "Optional ML module: model comparison, drivers and training history", 6.0)
shot("12_admin.png", "Administration: table sizes, schema objects, triggers, automated test results and audit trail", 6.0, 1500)
shot("15_mobile_dashboard.png", "Responsive layout at phone width (390 px)", 2.6)
L.heading(doc, "5.2 Integration with the DBMS", 2)
L.para(doc, "**Architecture.** Browser, then Flask routes (`app/app.py`), then the data-access layer (`app/db.py`, a python-oracledb connection pool in thin mode), then Oracle XEPDB1 "
            "with its stored procedures, package and views. The connection string, user and password come from environment variables with local defaults "
            "(`RETAINAI_DSN`, `RETAINAI_USER`, `RETAINAI_PASSWORD`).")
L.para(doc, "**Rules of the integration.** (1) The application changes data only by calling stored procedures, so validation and triggers always run. (2) Every query uses bind variables and never "
            "string concatenation, which closes the door on SQL injection. (3) A failed rule comes back as an ORA-20xxx message and is shown as a flash message. "
            "(4) Analytics come back as REF cursors and are converted to JSON for the charts. (5) Pages are plain views of the database: no business figure is calculated in Python, apart from the optional ML model.")
L.para(doc, "**Sign-in and password recovery.** The administrator password is checked against a hash. A laboratory setup has no mail server, so “Forgot password?” issues a six-digit one-time code and prints it "
            "in the terminal that runs the application. The code is valid for ten minutes and allows five attempts. A correct code lets the user choose a new password, which is stored as a salted hash in `app_config`; "
            "the change is also written to `audit_log`. Rebuilding the database restores the default login.")
L.table_caption(doc, "Table 5.1:", "Screens and the database objects they use")
L.table(doc, ["Screen / route", "Database objects"], [
    ["Dashboard `/`, `/api/*`", "pkg_retainai.dashboard_kpis, high_risk_report, churn_breakdown, segment_report; views v_risk_distribution, v_monthly_revenue"],
    ["Customers `/customers`, `/customers/new`, `/customers/<id>`", "v_customer_360; sp_add_customer, sp_cancel_subscription, sp_reactivate_subscription, sp_score_customer, fn_lifetime_value"],
    ["Billing `/billing`", "sp_generate_bills, sp_mark_overdue_bills, sp_record_payment"],
    ["Tickets `/tickets`", "sp_open_ticket, sp_close_ticket (trigger trg_tickets_ai)"],
    ["Retention `/retention`", "sp_create_campaign, sp_plan_retention, sp_record_outcome; view v_campaign_effectiveness"],
    ["Analytics and Reports `/analytics`, `/reports`", "pkg_retainai.*; views; CSV export of high_risk_report"],
    ["ML `/ml`", "reads v_customer_360; writes sp_save_ml_score, sp_log_ml_run; table ml_model_runs"],
    ["Admin `/admin`", "pkg_retainai.refresh_analytics, sp_set_reference_date; USER_OBJECTS, USER_TRIGGERS, test_results, audit_log"],
], [2.4, 3.8], size=9.5)
L.para(doc, "**Machine-learning module.** Training data are read from the view `v_customer_360`. Logistic Regression and Gradient Boosting are trained on a stratified 75/25 split, and the one with the better "
            "hold-out ROC-AUC is refitted on all customers; its probabilities are written back through `sp_save_ml_score`. In the final run Logistic Regression won "
            "(ROC-AUC %.3f, accuracy %.3f, precision %.3f, recall %.3f, F1 %.3f; " % tuple(LR) +
            "Gradient Boosting: ROC-AUC %.3f)" % GB[0] + ". Because the winner is refitted on every customer, the stored probabilities are in-sample: good for ranking, not for estimating accuracy.")

# =====================================================================================
# 6. VERIFICATION & TESTING
# =====================================================================================
L.heading(doc, "6. VERIFICATION & TESTING", 1, page_break=True)
L.heading(doc, "6.1 Test Cases", 2)
L.para(doc, f"Testing has two layers. The first is a PL/SQL script, `sql/12_tests.sql` (34 cases, TC01–TC34), which creates its own data through the public procedures, compares the results with independently "
            f"written SQL, records every outcome in the table `TEST_RESULTS` and removes its test data afterwards. The second (19 cases, UI01–UI19) drives the web application in a real browser. "
            f"The tables show the results as stored; the “Actual” column keeps only the first line of each message, because the long ORA-06512 stack lines add nothing. "
            f"**In the final run {n_sql_pass} of {n_sql} SQL tests and {n_ui_pass} of {n_ui} UI tests passed.**")
L.table_caption(doc, "Table 6.1:", "SQL and PL/SQL test cases (rows of TEST_RESULTS from the final run)")
L.table(doc, ["TC ID", "Test description", "Input", "Expected output", "Actual output", "Status"],
        [[t[0], t[1], t[2], t[3], t[4], t[5]] for t in tests], [0.45, 1.55, 1.15, 1.15, 1.35, 0.55], size=8, status_col=5)
L.table_caption(doc, "Table 6.2:", "Web front-end test cases (outputs/ui_tests.json from the final run)")
L.table(doc, ["TC ID", "Test description", "Input", "Expected output", "Actual output", "Status"],
        [[t["tc_id"], t["description"], t["input"], t["expected"], t["actual"].split(" | ")[0][:150], t["status"]] for t in ui],
        [0.45, 1.55, 1.0, 1.2, 1.45, 0.55], size=8, status_col=5)
L.heading(doc, "6.2 Output Screenshots", 2)
f6 = [0]
VS = os.path.join(EV, "vscode")


def shotV(fn, text, width=5.3):
    f6[0] += 1
    im = Image.open(os.path.join(VS, fn)); w, h = im.size
    im = im.crop((int(w * 0.278), int(h * 0.148), int(w * 0.767), int(h * 0.952)))
    g = im.convert("L"); rows = [y for y in range(im.height) if g.crop((0, y, im.width, y + 1)).getextrema()[1] > 110]
    if rows: im = im.crop((0, 0, im.width, min(im.height, rows[-1] + 60)))   # drop the empty part of the terminal below the output
    path = os.path.join(WORK, "v_" + fn); im.save(path)
    L.picture(doc, path, width); L.caption(doc, f"Figure 6.{f6[0]}: {text}")


L.para(doc, "The first seven figures are screenshots of the Visual Studio Code terminal. A short helper script ran the project's own commands in that terminal and saved an image of the VS Code window "
            "after each step, so what you see is the genuine terminal output; the helper only printed the grey step labels. Each image shows the end of the run, since a terminal cannot hold the whole output. "
            "The browser screenshots follow.")
shotV("01_build_finished.png", "python manage.py build finished: the last test cases of the SQL suite, all PASS")
shotV("02_objects.png", "Objects owned by the RETAINAI schema after the build, all VALID, with the main row counts")
shotV("03_functions.png", "The stored functions applied to six sample customers")
shotV("04_package.png", "Package pkg_retainai: KPI functions and a REF CURSOR printed in SQL*Plus")
shotV("05_cursors.png", "Cursor demonstrations (end of the output)")
shotV("06_triggers.png", "The ten triggers and the most recent audit-trail rows")
shotV("07_ui_tests.png", "python manage.py uitests: all 19 browser tests passed")


def shot6(fn, text, width=6.0, crop_h=None):
    f6[0] += 1
    im = Image.open(os.path.join(SH, fn)); path = os.path.join(WORK, "t_" + fn)
    if crop_h and im.height > crop_h: im = im.crop((0, 0, im.width, crop_h))
    im.save(path); L.picture(doc, path, width); L.caption(doc, f"Figure 6.{f6[0]}: {text}")


shot6("17_customer_registered.png", "UI04 – customer registered through the form; the new customer is already scored", 5.6)
shot6("18_duplicate_email_error.png", "UI05 – duplicate e-mail rejected by the database (ORA-20002 shown to the user)", 5.6)
shot6("19_ticket_opened.png", "UI07 – urgent cancellation ticket opened; a retention call is queued by the trigger", 5.6)
shot6("20_subscription_cancelled.png", "UI08 – subscription cancelled; customer CHURNED and a win-back action created", 5.6)
shot6("21_billing_run.png", "UI10 – billing run is idempotent (second run creates 0 bills)", 5.6)
shot6("22_retention_after.png", "UI12–UI14 – campaign created, actions planned, outcome recorded", 5.6, 1500)
shot6("25_reset_wrong_code.png", "UI17 – a wrong recovery code is rejected", 5.6)
f6[0] += 1
COMPILE_FIG = "6.%d" % f6[0]
L.picture(doc, spool_image("verify.txt", "=== V3.", "=== V4."), 4.8)
L.caption(doc, "Figure " + COMPILE_FIG + ": Compilation check – no invalid objects and no compilation errors after the build (real output)")
L.heading(doc, "6.3 Validation of Functionality", 2)
L.para(doc, "The results support five conclusions. **Constraints** reject bad data: an invalid e-mail (TC04), a negative charge (TC05), an unknown customer (TC06) and a second active subscription (TC07). "
            "**Triggers** do their job: normalisation (TC02), blocked deletion (TC08), the cascade from cancellation to CHURNED plus a win-back action (TC13), the audit trail (TC15), the single "
            "compound-trigger audit row (TC17), unchangeable paid bills (TC19–TC20) and the automatic retention call (TC23). **Procedures** are correct and repeatable: the billing run creates "+BILLS_N+" bills "
            "the first time and none the second time (TC16), and an accepted outcome applies its discount exactly once (TC33). **Analytics** agree with independent SQL (TC25–TC27). "
            "The **front end** performs the same operations through the browser (UI01–UI19), including the password-recovery flow (UI17–UI19).")
L.para(doc, "**Defects found by running the code.** The first complete run on Oracle exposed three defects that no amount of reading had caught, and all three were fixed before the final run. "
            "`ORA_HASH` is a SQL function and cannot be used in a PL/SQL expression, which gave ORA-06550/PLS-00201 in the seeding script; it was replaced by `DBMS_UTILITY.GET_HASH_VALUE`. "
            "A `BOOLEAN` parameter had been used inside an SQL `CASE` in the test logger (PLS-00905), so the value is now worked out in a PL/SQL variable first. "
            "Test TC17 failed because the earlier demonstration script had left its own BULK_INSERT audit row and the test counted two; it now counts only rows created after it started, and the trigger itself was correct. "
            "A fourth problem appeared when the build was repeated for the final screenshots: `01_schema.sql` fails on tables that already exist, and the seed script then added another "+TICKETS_N+" tickets "
            "on top of the old ones, so the counts grew with every run. `sql/00_reset.sql` now drops the previous objects at the start of `run_all.sql`, which makes the build repeatable.")
L.para(doc, "**Limits of the testing.** One person ran the tests on one machine with one browser. There was no load, concurrency or security testing, and the web tests rely on the data the build creates. "
            "The test data are deleted afterwards, but their audit-trail rows are kept as evidence, which is why the audit log holds several hundred rows. "
            "The recovery-code test reads the code from the application's own log file, because that is where the code is printed.")

# =====================================================================================
# 7. RESULTS & DISCUSSION
# =====================================================================================
L.heading(doc, "7. RESULTS & DISCUSSION", 1, page_break=True)
L.heading(doc, "7.1 Achievements", 2)
L.bullets(doc, [
    "Objective 1: a 17-table 3NF schema now holds 7,043 real customers, and the loaded churn count (1,869) equals the source (Figure 3.4).",
    "Objective 2: 7 functions, 20 procedures, 1 package, 10 triggers and all five cursor types compile without errors (Figure "+COMPILE_FIG+") and pass 34 automated tests.",
    "Objective 3: eleven reports and eight views; the data confirm the usual churn drivers (month-to-month 42.7%, fibre optic 41.9%, electronic check 45.3%, first-year customers 47.4%).",
    f"Objective 4: a responsive web application that performs every operation through stored procedures, with {n_ui_pass} passing browser tests and {N_SHOTS} real screenshots.",
    "Objective 5: the ML module reaches ROC-AUC %.3f on a hold-out set" % LR[0] + ", in line with published results for this dataset, and writes its probabilities back to Oracle.",
    "Objective 6: every result in this report is the recorded output of a real run.",
])
L.para(doc, "The "+HIGH_N+" active HIGH-risk customers account for "+RAR_TXT+" of monthly revenue at risk ("+RAR_PCT+"% of MRR). The demonstration campaigns reached an acceptance rate of 26.7%–30.9%, "
            "but the responses are simulated, so that figure only shows the reporting works; it says nothing about how real offers would perform.")
L.heading(doc, "7.2 Challenges Faced", 2)
L.bullets(doc, [
    "**Restricted pluggable database.** On the author’s laptop XEPDB1 was open in RESTRICTED mode because of an unrelated pending synchronisation error (ORA-65177 during a `GRANT DBA TO login CONTAINER = ALL`, recorded in `PDB_PLUG_IN_VIOLATIONS`), so an ordinary user could not connect (ORA-01035). "
    "A non-destructive re-synchronisation did not clear it, so the project user was granted `RESTRICTED SESSION`; this is documented in `sql/00_create_user.sql`. The unrelated common user was left alone.",
    "**No calendar in the dataset.** The Telco file has tenure but no dates. A fixed business reference date (30-Sep-2026, table `app_config`) makes every date rule reproducible.",
    "**Mutating-table risk.** Cancelling a subscription updates the customer and inserts an action, so the triggers were written so that none of them reads the table it is changing. The compound trigger collects row values in memory and writes once after the statement.",
    "**Synthetic data.** Customer names, billing history, tickets and campaign responses had to be generated. They are labelled everywhere, and the built-in link between tickets and churn is acknowledged.",
    "**Slow first requests.** On this laptop the first dashboard request after a rebuild was slow, so the browser tests use long time-outs. Nothing beyond indexes and statistics was tuned, and the cause was not investigated.",
    "**Evidence.** Long SQL*Plus outputs are shown as images rendered from the real spool files because the text does not fit on a screen. The terminal screenshots in Section 6.2 show the end of the same runs.",
    "**No mail server.** Password recovery therefore prints its one-time code in the terminal instead of sending an e-mail.",
])
L.heading(doc, "7.3 Possible Improvements", 2)
L.bullets(doc, [
    "Performance: partition `billing_records` by month, add materialised views for the dashboard, cache the KPI results and look into the slow first request.",
    "Security: a separate schema owner and application user with minimum privileges, database roles for agents and managers, per-user login in the web layer, and Oracle Virtual Private Database for row-level access.",
    "Usability: pagination and server-side search for large lists, editing of campaigns and tickets, and notifications for agents.",
    "Analytics: retrain the ML model on a schedule with a train/validation split stored per customer, and add a survival analysis of tenure.",
])

# =====================================================================================
# 8. CONCLUSION
# =====================================================================================
L.heading(doc, "8. CONCLUSION", 1, page_break=True)
L.heading(doc, "8.1 Summary of Learning Outcomes", 2)
L.para(doc, "The project pulled most of the database syllabus into one realistic problem. On the design side that meant ER modelling, a lossless decomposition to 3NF/BCNF and DDL with every kind of constraint, "
            "including a function-based unique index. On the programming side it meant ETL with `INSERT ... SELECT` and analytic functions; PL/SQL functions and procedures with exception handling and transaction control; "
            "packages with REF cursors; row-level, statement-level and compound triggers; autonomous transactions; implicit, explicit, parameterised, bulk-collect and REF cursors; and advanced SQL such as window functions, "
            "ROLLUP and anti-joins. The web application showed how an application should talk to Oracle: through bind variables and stored procedures, and nothing else. "
            "The most useful lesson was about testing. Running the code found defects that reading it never did.")
L.heading(doc, "8.2 Future Scope", 2)
L.para(doc, "The same schema could run on a cloud Oracle service with role-based access for several users and take live data from a billing system. The planned retention actions could go out by e-mail or SMS, "
            "a database scheduler job could refresh the ML scores every night, and a REST layer over the existing stored procedures would let a mobile app reuse the same logic.")

# =====================================================================================
# 9. REFERENCES
# =====================================================================================
L.heading(doc, "9. REFERENCES", 1, page_break=True)
refs = [
    "A. Silberschatz, H. F. Korth, and S. Sudarshan, Database System Concepts, 7th ed. New York, NY, USA: McGraw-Hill, 2019.",
    "R. Elmasri and S. B. Navathe, Fundamentals of Database Systems, 7th ed. Boston, MA, USA: Pearson, 2016.",
    "Oracle Corporation, “PL/SQL Language Reference, Oracle Database 21c,” Oracle Help Center. [Online]. Available: https://docs.oracle.com/en/database/oracle/oracle-database/21/lnpls/",
    "Oracle Corporation, “SQL Language Reference, Oracle Database 21c,” Oracle Help Center. [Online]. Available: https://docs.oracle.com/en/database/oracle/oracle-database/21/sqlrf/",
    "Oracle Corporation, “SQL*Plus User’s Guide and Reference, 21c,” Oracle Help Center. [Online]. Available: https://docs.oracle.com/en/database/oracle/oracle-database/21/sqpug/",
    "BlastChar, “Telco Customer Churn (IBM sample data set),” Kaggle. [Online]. Available: https://www.kaggle.com/datasets/blastchar/telco-customer-churn",
    "Oracle Corporation, “python-oracledb documentation.” [Online]. Available: https://python-oracledb.readthedocs.io/",
    "Pallets Projects, “Flask documentation.” [Online]. Available: https://flask.palletsprojects.com/",
    "F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” J. Mach. Learn. Res., vol. 12, pp. 2825–2830, 2011.",
    "Microsoft, “Playwright for Python.” [Online]. Available: https://playwright.dev/python/",
]
for i, r in enumerate(refs, 1):
    p = L.para(doc, f"[{i}] {r}", align="left", after=6)
    p.paragraph_format.left_indent = Inches(0.45); p.paragraph_format.first_line_indent = Inches(-0.45)
L.para(doc, "Web resources were retrieved by the author in October 2026.", align="left", italic=True, size=11)

# =====================================================================================
# 10. APPENDICES
# =====================================================================================
L.heading(doc, "10. APPENDICES", 1, page_break=True)
L.heading(doc, "Appendix A – Additional Code", 2)
L.para(doc, "The code that did not fit in the main body. The data-load script `sql/02_load_staging.sql` (7,043 INSERT statements, generated by `scripts/gen_staging.py`) is not printed, "
            "and the templates and style sheets stay in the project folder.")
APP_FILES = [
    ("A.1 sql/00_create_user.sql – schema owner", os.path.join(SQL, "00_create_user.sql")),
    ("A.1b sql/00_reset.sql – clean start for a repeated build", os.path.join(SQL, "00_reset.sql")),
    ("A.2 sql/04_etl.sql – ETL procedure", os.path.join(SQL, "04_etl.sql")),
    ("A.3 sql/05_views.sql – views", os.path.join(SQL, "05_views.sql")),
    ("A.4 sql/06_procedures.sql – all procedures", os.path.join(SQL, "06_procedures.sql")),
    ("A.5 sql/09_seed_ops.sql – seed operations", os.path.join(SQL, "09_seed_ops.sql")),
    ("A.6 sql/11_reports.sql – analytical reports", os.path.join(SQL, "11_reports.sql")),
    ("A.7 sql/12_tests.sql – automated SQL test cases", os.path.join(SQL, "12_tests.sql")),
    ("A.8 sql/13_verify.sql – verification queries", os.path.join(SQL, "13_verify.sql")),
    ("A.9 sql/14_demo_routines.sql – demonstration of functions, procedures and triggers", os.path.join(SQL, "14_demo_routines.sql")),
    ("A.10 sql/run_all.sql – master script", os.path.join(SQL, "run_all.sql")),
    ("A.11 app/db.py – Oracle access layer", os.path.join(APP, "db.py")),
    ("A.12 app/ml.py – machine-learning module", os.path.join(APP, "ml.py")),
    ("A.13 app/app.py – Flask application", os.path.join(APP, "app.py")),
    ("A.14a manage.py \u2013 project command runner", os.path.join(PROJ, "manage.py")),
    ("A.14 scripts/capture.py – browser tests and screenshots", os.path.join(PROJ, "scripts", "capture.py")),
]
for title_, path in APP_FILES:
    L.subheading(doc, title_)
    L.code_block(doc, rd(path), size=7.5)

L.heading(doc, "Appendix B – Additional Screenshots", 2)
fB = [0]


def shotB(fn, text, width=5.6, crop_h=None):
    fB[0] += 1
    im = Image.open(os.path.join(SH, fn)); path = os.path.join(WORK, "b_" + fn)
    if crop_h and im.height > crop_h: im = im.crop((0, 0, im.width, crop_h))
    im.save(path); L.picture(doc, path, width); L.caption(doc, f"Figure B.{fB[0]}: {text}")


shotB("04_customers_high_risk.png", "Customer list filtered to ACTIVE customers with HIGH risk, sorted by risk", 5.6, 1300)
shotB("14_ml_prediction.png", "ML what-if prediction for a hypothetical customer", 5.6)
shotB("16_mobile_customers.png", "Customer list at phone width", 2.6)
shotB("23_admin_after.png", "Administration page after the browser tests (audit entries written by triggers)", 5.6, 1500)
shotB("22_retention_after.png", "Retention page after planning actions and recording outcomes (top of page)", 5.6, 1600)
L.para(doc, "All screenshots and the logs of the final run are in the folder `outputs/` of the project.")

L.heading(doc, "Appendix C \u2013 How to Run the Project", 2)
L.para(doc, "Needs Oracle Database 21c XE, Python 3.10 or later and Visual Studio Code with the Python extension. Open the project folder in VS Code and run these in the integrated terminal:")
L.code_block(doc, "python manage.py setup      # install packages, create the RETAINAI schema owner, build the database, run the 34 SQL tests\npython manage.py run        # web application at http://127.0.0.1:5000  (user admin, password retainai)\npython manage.py uitests    # 19 browser test cases and screenshots (written to outputs/)", size=9)
doc.save(OUT)
print("saved", OUT, "abstract words:", wc)
