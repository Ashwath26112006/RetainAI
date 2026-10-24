"""Helpers used by build_report.py (python-docx primitives that follow the BCSE302P template conventions)."""
import re, copy
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFont

TNR = "Times New Roman"
CODE = "Courier New"
AL = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
      "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}


def _fonts(run, name=TNR, size=12, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)


def _spacing(p, after=6, before=0, line=1.5):
    pf = p.paragraph_format
    pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = line


_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|_\_.+?_\_)")


def add_runs(p, text, size=12, bold=False, italic=False, color=None):
    """Mini markup: **bold**, `code`."""
    for part in _TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); _fonts(r, TNR, size, True, italic, color)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); _fonts(r, CODE, size - 1.5, bold, italic, color)
        else:
            r = p.add_run(part); _fonts(r, TNR, size, bold, italic, color)


def para(doc, text, align="both", after=6, before=0, size=12, bold=False, italic=False, keep_next=False, color=None, indent=None):
    p = doc.add_paragraph()
    p.alignment = AL[align]
    _spacing(p, after, before)
    if keep_next: p.paragraph_format.keep_with_next = True
    if indent is not None: p.paragraph_format.left_indent = Inches(indent)
    add_runs(p, text, size, bold, italic, color)
    return p


def bullets(doc, items, size=12, after=3):
    for it in items:
        p = doc.add_paragraph()
        p.alignment = AL["both"]
        _spacing(p, after)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.35); pf.first_line_indent = Inches(-0.2)
        r = p.add_run("•  "); _fonts(r, TNR, size)
        add_runs(p, it, size)


def numbered(doc, items, size=12, after=3):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.alignment = AL["both"]
        _spacing(p, after)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.4); pf.first_line_indent = Inches(-0.3)
        r = p.add_run(f"{i}.  "); _fonts(r, TNR, size)
        add_runs(p, it, size)


def heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr(); ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), "Heading%d" % level); ppr.insert(0, ps)
    pf = p.paragraph_format
    pf.space_after = Pt(10 if level == 1 else 6); pf.space_before = Pt(6 if level == 1 else 10)
    pf.line_spacing = 1.5
    pf.keep_with_next = True
    if page_break: pf.page_break_before = True
    r = p.add_run(text); _fonts(r, TNR, 14, True, None, "000000")
    return p


def subheading(doc, text, size=12):
    p = doc.add_paragraph(); _spacing(p, 4, 6); p.paragraph_format.keep_with_next = True
    r = p.add_run(text); _fonts(r, TNR, size, True)
    return p


def shade(p, fill="F2F2F2"):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def code_block(doc, text, size=8.5, max_chars=None):
    lines = text.expandtabs(2).rstrip("\n").split("\n")
    for ln in lines:
        p = doc.add_paragraph()
        p.alignment = AL["left"]
        _spacing(p, 0, 0, 1.0)
        p.paragraph_format.keep_together = True
        r = p.add_run(ln if ln else " "); _fonts(r, CODE, size)
        shade(p)
    sp = doc.add_paragraph(); _spacing(sp, 4, 0, 1.0)


def caption(doc, text, after=10):
    p = doc.add_paragraph(); p.alignment = AL["center"]; _spacing(p, after, 2)
    r = p.add_run(text); _fonts(r, TNR, 11, None, True)
    return p


def table_caption(doc, label, text):
    p = doc.add_paragraph(); p.alignment = AL["left"]; _spacing(p, 4, 6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(label + " "); _fonts(r, TNR, 12, True)
    r2 = p.add_run(text); _fonts(r2, TNR, 12)
    return p


def picture(doc, path, width_in=6.2, crop=None, keep_next=True):
    p = doc.add_paragraph(); p.alignment = AL["center"]; _spacing(p, 2, 4, 1.0)
    if keep_next: p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width_in))
    return p


def _borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement("w:" + e); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), "000000"); b.append(x)
    tblPr.append(b)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)


def _cell(cell, text, size, bold=False, fill=None, align="left", font=TNR):
    cell.text = ""
    p = cell.paragraphs[0]; p.alignment = AL[align]; _spacing(p, 0, 0, 1.0)
    for i, line in enumerate(str(text).split("\n")):
        if i:
            p = cell.add_paragraph(); p.alignment = AL[align]; _spacing(p, 0, 0, 1.0)
        r = p.add_run(line); _fonts(r, font, size, bold)
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", 40), ("left", 70), ("bottom", 40), ("right", 70)):
        e = OxmlElement("w:" + side); e.set(qn("w:w"), str(w)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)
    if fill:
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcPr.append(shd)


def table(doc, header, rows, widths, size=10, align=None, code_cols=(), status_col=None):
    t = doc.add_table(rows=1, cols=len(header))
    _borders(t)
    t.autofit = False
    align = align or ["left"] * len(header)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; _cell(c, h, size, True, "D9E2F3", "left")
        c.width = Inches(widths[i])
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            fill = None
            if status_col is not None and i == status_col:
                fill = "E2F0D9" if str(v).startswith("PASS") else ("F8D7DA" if str(v).startswith("FAIL") else None)
            _cell(cells[i], v, size, False, fill, align[i], CODE if i in code_cols else TNR)
            cells[i].width = Inches(widths[i])
        trPr = t.rows[-1]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    sp = doc.add_paragraph(); _spacing(sp, 4, 0, 1.0)
    return t


# ------------------------------------------------------------------ images
def _font(size):
    for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",):
        try:
            return ImageFont.truetype(f, size)
        except Exception:
            pass
    return ImageFont.load_default()


def render_text(lines, out, header, size=15):
    """Render real spool text as a dark 'console-style' picture. The header states that it is a rendering."""
    f = _font(size); hb = _font(size - 2)
    cw = f.getlength("M")
    lh = int(size * 1.35)
    width = int(max(len(l) for l in lines + [header]) * cw) + 36
    width = max(width, 760)
    height = lh * len(lines) + 62
    img = Image.new("RGB", (width, height), "#0f172a")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, width, 30], fill="#1e293b")
    d.text((14, 7), header, font=hb, fill="#94a3b8")
    y = 44
    for l in lines:
        col = "#e2e8f0"
        if l.startswith("===") or l.startswith("---") and "=" not in l[:3] and len(l) > 3 and l[3] == " ":
            col = "#5eead4"
        if re.match(r"^(ORA-|PLS-|SP2-|ERROR)", l):
            col = "#fca5a5"
        d.text((18, y), l, font=f, fill=col)
        y += lh
    img.save(out)
    return out


def squeeze(lines):
    out, blank = [], 0
    for l in lines:
        l = l.rstrip()
        if not l:
            blank += 1
            if blank > 1:
                continue
        else:
            blank = 0
        out.append(l)
    while out and not out[0]: out.pop(0)
    while out and not out[-1]: out.pop()
    return out
